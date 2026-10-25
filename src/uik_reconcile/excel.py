"""Импорт с внутренней копией; исходный файл никогда не изменяется."""
import os
import shutil
import uuid
from pathlib import Path

from openpyxl import load_workbook


def sheet_names(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_table(path, sheet):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book[sheet].iter_rows(values_only=True))
        if not rows or not any(v is not None for v in rows[0]):
            return [], []
        width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
        headers = [str(v).strip() if v is not None else "" for v in rows[0][:width]]
        validate_headers(headers)
        data = [dict(zip(headers, list(row[:width]) + [None] * max(0, width-len(row))))
                for row in rows[1:] if any(v is not None for v in row[:width])]
        return headers, data
    finally:
        book.close()


def validate_headers(headers):
    if any(not h for h in headers):
        raise ValueError("В первой строке есть пустые заголовки. Укажите названия всех столбцов.")
    if len(set(headers)) != len(headers):
        raise ValueError("В первой строке есть повторяющиеся заголовки")


def copy_internal(root, source):
    dest = Path(root) / "imports" / (uuid.uuid4().hex + ".xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    temp = dest.with_suffix(".tmp")
    try:
        shutil.copy2(source, temp)
        os.replace(temp, dest)
    finally:
        temp.unlink(missing_ok=True)
    return dest.relative_to(root).as_posix()
